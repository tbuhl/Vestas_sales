"""Extract the installed-base track record (per country and per turbine type) from the Vestas workbook.

Usage:
    python tools/build_installed_base.py path/to/Vestas_economical_data.xlsx

Writes data/vestas_installed_base.json, which the dashboard reads when no workbook is present.
The workbook sheets used are `Track_record_country_end_23` and `Track record turbines`.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

OUT_FILE = Path(__file__).resolve().parent.parent / "data" / "vestas_installed_base.json"
COUNTRY_SHEET = "Track_record_country_end_23"
TURBINE_SHEET = "Track record turbines"


def _num(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_country_rows(ws) -> list[dict]:
    rows = []
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0 or not row or not row[0]:
            continue
        name = str(row[0]).strip()
        qty, mw, last = _num(row[1]), _num(row[2]), _num(row[3])
        if not name or mw is None:
            continue
        rows.append(
            {
                "country": name,
                "turbines": int(qty) if qty is not None else None,
                "installed_mw": mw,
                "latest_installation_year": int(last) if last is not None else None,
            }
        )
    return rows


def read_turbine_rows(ws) -> tuple[list[dict], dict | None]:
    """Return (turbine type rows, the workbook's own `Total` row kept separate to avoid double counting)."""
    rows, total = [], None
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0 or not row or not row[0]:
            continue
        name = str(row[0]).strip()
        qty, mw = _num(row[1]), _num(row[2])
        if not name or mw is None:
            continue
        entry = {"turbine_type": name, "turbines": int(qty) if qty is not None else None, "installed_mw": mw}
        if name.lower() == "total":
            total = entry
        else:
            rows.append(entry)
    return rows, total


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    workbook = Path(sys.argv[1])
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    turbine_rows, turbine_total = read_turbine_rows(wb[TURBINE_SHEET])
    payload = {
        "source": workbook.name,
        "as_of_year": 2023,
        "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "countries": read_country_rows(wb[COUNTRY_SHEET]),
        "turbine_types": turbine_rows,
        "turbine_total": turbine_total,
    }
    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUT_FILE.write_text(json.dumps(payload, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {OUT_FILE} ({len(payload['countries'])} countries, {len(payload['turbine_types'])} turbine types)")


if __name__ == "__main__":
    main()
